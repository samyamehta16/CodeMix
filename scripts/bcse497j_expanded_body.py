"""Expanded chapter body for BCSE497J TruthGuard report."""
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

IMAGES = Path(__file__).resolve().parents[1] / "images"


def _font(run, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_body(doc, text, *, size=12, space_after=8, first_indent=True):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.15
    if first_indent:
        para.paragraph_format.first_line_indent = Inches(0.3)
    run = para.add_run(text)
    _font(run, size=size)
    return para


def add_heading_line(doc, text, *, size=14, bold=True, center=False, space_before=14):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    _font(run, size=size, bold=bold)
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(f"• {text}")
    _font(run, size=12)
    return para


def add_caption(doc, text, *, above=False):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4 if not above else 10)
    para.paragraph_format.space_after = Pt(10 if not above else 4)
    run = para.add_run(text)
    _font(run, size=10, italic=True)
    return para


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _font(run, size=10, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            _font(run, size=10)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_picture(doc, path: Path, width=5.6):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(path), width=Inches(width))
    return para


def add_full_body(doc):
    # ------------------------------------------------------------------
    # ABSTRACT (~450 words)
    # ------------------------------------------------------------------
    add_heading_line(doc, "ABSTRACT", size=16, bold=True, center=True, space_before=0)
    add_body(
        doc,
        "Misinformation on Tamil social media is difficult to detect automatically for two "
        "practical reasons. First, users frequently write in Tanglish — Tamil and English mixed "
        "inside the same sentence, often with Romanized Tamil spellings such as “pannitanga” or "
        "“solli irukanga.” Many off-the-shelf detectors assume either clean English or formal "
        "native-script Tamil, so they fragment or misread such input. Second, a supervised "
        "classifier trained only on yesterday’s headlines cannot verify a rumour that started "
        "today unless it can retrieve fresh evidence from the web. Style-based scoring alone is "
        "not enough when the truth depends on current announcements from institutions such as "
        "the RBI, ISRO or a district collector.",
        first_indent=False,
    )
    add_body(
        doc,
        "This B.Tech. Project-I (BCSE497J) work presents TruthGuard, a dual-track system. "
        "Approach A is the supervised classification track. We collected and cleaned a balanced "
        "Tamil veracity corpus of 1400 headlines (exactly 700 FAKE and 700 REAL) from Tamil "
        "fact-checking publishers and mainstream news sites. We fine-tuned IndicBERTv2, MuRIL, "
        "mBERT and XLM-RoBERTa under a common protocol, and compared them with classical TF-IDF "
        "baselines (Naive Bayes, Logistic Regression, LinearSVC) plus a style-only logistic "
        "baseline. Because FAKE rows come almost only from fact-checkers and REAL rows from news "
        "outlets, a random train/test split can drive accuracy near 100% by learning source or "
        "style shortcuts rather than genuine veracity. Therefore our primary evaluation uses a "
        "source-disjoint hard split: training publishers are held out from test publishers. On "
        "that hard test (382 headlines), IndicBERTv2 reached about 86.65% accuracy with macro-F1 "
        "about 0.623. MuRIL and mBERT stayed near 61%. XLM-RoBERTa collapsed to about 32% accuracy "
        "and predicted FAKE for a large share of REAL headlines. TF-IDF models remained competitive "
        "on accuracy, which is an important honesty check for a student project.",
        first_indent=False,
    )
    add_body(
        doc,
        "Approach B is the evidence-grounded track for Tanglish claims. It runs six stages: "
        "claim extraction, query generation, Serper web search, webpage text extraction "
        "(Trafilatura / BeautifulSoup, with search-snippet fallback), multi-factor ranking "
        "(relevance, source authority, freshness), and Gemini adjudication that must cite "
        "evidence IDs. Modules that search or extract never decide truth. When evidence is "
        "missing, the system prefers UNVERIFIED rather than FAKE. A Streamlit demo shows live "
        "runs on prepared claims (FAKE currency rumour, REAL Chandrayaan landing, MISLEADING "
        "holiday wording, UNVERIFIED free-iPhone forward). Approach A is fast for offline "
        "screening of formal Tamil headlines; Approach B is slower and API-dependent but can "
        "check new rumours with citations. Remaining limits include source–label alignment in "
        "the 1400 set, class imbalance on the hard test FAKE subset, and the absence of a large "
        "labelled Tanglish veracity benchmark for quantitative Approach B metrics.",
        first_indent=False,
    )
    add_body(
        doc,
        "Keywords: Fake news detection; Tanglish; code-mixed NLP; MuRIL; IndicBERTv2; mBERT; "
        "XLM-RoBERTa; LLM; evidence-grounded verification; Serper; Streamlit; source-disjoint evaluation.",
        first_indent=False,
    )

    # ------------------------------------------------------------------
    # CHAPTER 1
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 1", size=16, bold=True, center=True)
    add_heading_line(doc, "INTRODUCTION", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "1.1 Background", size=12, bold=True, space_before=10)
    add_body(
        doc,
        "Mobile internet and messaging apps have made it easy to share news-like posts in Tamil "
        "Nadu. Alongside genuine updates, false or half-true claims also travel quickly — for "
        "example rumours about new currency notes, free gadgets from “government schemes,” or "
        "statewide school holidays after rain. When such claims are wrong, they can create panic, "
        "waste public attention, or damage trust in institutions.",
    )
    add_body(
        doc,
        "A second difficulty is linguistic. Everyday Tamil social text is often code-mixed. Users "
        "mix Tamil script, Romanized Tamil, and English inside one utterance. An example used "
        "throughout this project is: “RBI 5000 rupees note release pannitanga nu news la varudhu, "
        "real ah fake ah verify pannunga.” Tools built only for monolingual English fake-news "
        "corpora, or only for edited newspaper Tamil, do not see this pattern during training. "
        "Tokenizers may split Romanized Tamil oddly, and lexical cues differ from formal headlines.",
    )
    add_body(
        doc,
        "There is also a knowledge problem. Even a strong classifier freezes its notion of the "
        "world at training time. It cannot “look up” whether RBI announced a note today. "
        "Evidence-based fact-checking systems therefore retrieve documents first and decide later. "
        "TruthGuard uses both ideas: a trainable classifier for formal Tamil headlines, and a "
        "retrieve-then-verify pipeline for Tanglish claims.",
    )

    add_heading_line(doc, "1.2 Motivations", size=12, bold=True)
    add_body(
        doc,
        "Most published high scores in fake-news detection come from English datasets. Those "
        "numbers are useful as references, but they are not directly transferable to Tanglish "
        "WhatsApp forwards. We also noticed that if FAKE and REAL examples come from different "
        "website families, a model can quietly learn “this looks like a fact-check page” versus "
        "“this looks like a news page,” which inflates accuracy on a random split.",
    )
    add_body(
        doc,
        "Our motivation was therefore threefold: (1) build a self-collected Tamil veracity set "
        "we fully control and can document; (2) evaluate with a source-disjoint protocol so Review "
        "panels see honest numbers; (3) implement a working Tanglish demo that shows citations, "
        "not only a label. The project sits under BCSE497J Project-I and is meant to be "
        "reproducible with free-tier APIs and CPU-friendly model sizes.",
    )

    add_heading_line(doc, "1.3 Scope of the Project", size=12, bold=True)
    add_body(doc, "In scope for this Project-I phase:", first_indent=False)
    add_bullet(doc, "Approach A: binary REAL/FAKE classification on Tamil headlines from the 1400 balanced set.")
    add_bullet(doc, "Random split (diagnostic) and hard source-disjoint split (primary) evaluation.")
    add_bullet(doc, "Comparison across IndicBERTv2, MuRIL, mBERT, XLM-RoBERTa and TF-IDF baselines.")
    add_bullet(doc, "Approach B: six-stage evidence pipeline for Tanglish claims with Streamlit UI.")
    add_bullet(doc, "Qualitative demo cases covering FAKE, REAL, MISLEADING and UNVERIFIED outcomes.")
    add_body(doc, "Out of scope / deferred:", first_indent=False)
    add_bullet(doc, "Production deployment, mobile app packaging, or large-scale moderation APIs.")
    add_bullet(doc, "Full Acc/F1 for Approach B on a large gold Tanglish veracity benchmark (not yet available).")
    add_bullet(doc, "Training very large LLMs from scratch or heavy multi-GPU hyperparameter sweeps.")
    add_bullet(doc, "Legal takedown workflows; the system is a research/demo aid, not a censor.")

    # ------------------------------------------------------------------
    # CHAPTER 2
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 2", size=16, bold=True, center=True)
    add_heading_line(doc, "PROJECT DESCRIPTION AND GOALS", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "2.1 Literature Review", size=12, bold=True)
    add_body(
        doc,
        "Classical fake-news detection used handcrafted features: n-grams, punctuation, sentiment "
        "scores, readability measures and sometimes user or network metadata. These methods are "
        "interpretable but brittle across domains. Deep learning then moved the field toward "
        "end-to-end text encoders. BERT-based systems such as FakeBERT combine transformer "
        "embeddings with CNN blocks and report very high accuracy on English Kaggle-style corpora. "
        "Those results show that transformers can model news text well in English, but the "
        "evaluation setting is monolingual and closed-world: the model never retrieves external "
        "evidence for a new claim.",
    )
    add_body(
        doc,
        "For Indian languages, general multilingual models (mBERT, XLM-RoBERTa) help, yet capacity "
        "for Indic scripts and transliteration varies. MuRIL was pretrained on 17 Indian languages "
        "plus English with translation/transliteration signal, which is relevant for Romanized "
        "Tamil. IndicBERT and IndicBERTv2 from AI4Bharat target Indic corpora more directly; we "
        "use IndicBERTv2 because the older ALBERT-based IndicBERT v1 checkpoint is gated on "
        "Hugging Face. Shared tasks under DravidianLangTech / related workshops created Tanglish "
        "resources mainly for sentiment and offensive language, not for evidence-grounded "
        "fake-news verification.",
    )
    add_body(
        doc,
        "Evidence-based fact checking is a separate line of work. FEVER asks systems to retrieve "
        "Wikipedia sentences and label claims as supported, refuted or not enough info. "
        "Retrieval-augmented generation (RAG) couples a retriever with a generator so answers "
        "are conditioned on documents. Our Approach B follows the same separation: search and "
        "extraction never assign REAL/FAKE; only the final verifier does, and only from the "
        "evidence text it was given. We specialize queries and authority lists toward Indian "
        "sources (RBI, PIB, major fact-checkers, reputable news).",
    )

    add_heading_line(doc, "2.2 Research Gap", size=12, bold=True)
    add_body(
        doc,
        "After reviewing the above, three gaps remain for a Tamil/Tanglish student project. "
        "First, many Tamil experiments reuse external dumps without documenting source–label "
        "alignment, so near-perfect accuracy may be leakage. Second, Tanglish code-mixing is "
        "studied more for sentiment/offense than for veracity with citations. Third, few Project-I "
        "systems show both a measurable Indic classifier track and a live retrieve-then-verify "
        "demo in one codebase. TruthGuard addresses these by (a) releasing evaluation on a "
        "self-collected 1400 set with an explicit hard split, and (b) shipping Approach B with "
        "Streamlit screenshots and module-level separation.",
    )

    add_heading_line(doc, "2.3 Objectives", size=12, bold=True)
    add_bullet(doc, "Collect, clean and document a balanced 1400-row Tamil FAKE/REAL headline dataset with source metadata.")
    add_bullet(doc, "Implement random and source-disjoint splits; treat hard split as the primary metric table.")
    add_bullet(doc, "Fine-tune and compare IndicBERTv2, MuRIL, mBERT and XLM-RoBERTa under one training protocol.")
    add_bullet(doc, "Add TF-IDF and style-only baselines to contextualize transformer scores.")
    add_bullet(doc, "Build a six-stage Approach B pipeline with caching, failure notes and citation constraints.")
    add_bullet(doc, "Demonstrate Approach B in Streamlit on FAKE/REAL/MISLEADING/UNVERIFIED Tanglish presets.")
    add_bullet(doc, "State limitations clearly for Review panels (confounds, API limits, future work).")

    add_heading_line(doc, "2.4 Problem Statement", size=12, bold=True)
    add_body(
        doc,
        "Approach A (closed-world classification). Given a formal Tamil headline x_formal, learn a "
        "function f_A such that f_A(x_formal) ∈ {REAL, FAKE}. Training uses labelled headlines from "
        "the 1400 corpus. Evaluation must disclose whether train and test share publishers.",
    )
    add_body(
        doc,
        "Approach B (open-world verification). Given informal Tanglish text x_cm, (i) detect whether "
        "a checkable claim exists, (ii) normalize it to a clear English proposition c, (iii) retrieve "
        "an evidence set E = {e1, …, ek} from the live web, and (iv) produce "
        "v = f_B(c, E) ∈ {REAL, FAKE, MISLEADING, UNVERIFIED} together with cited evidence IDs and a "
        "short explanation grounded in those sources. If E is empty or insufficient, return UNVERIFIED.",
    )

    add_heading_line(doc, "2.5 Project Plan", size=12, bold=True)
    add_caption(doc, "Table 2.1 High-level project phases", above=True)
    add_table(
        doc,
        ["Phase", "Work items", "Outcome"],
        [
            ["1. Data & scaffolding", "Collect 1400 set; clean; earlier LID silver labels on YouTube comments", "Documented corpus + tooling"],
            ["2. Approach A", "Preprocess splits; train transformers + TF-IDF; save metrics", "Hard/random result tables"],
            ["3. Approach B", "Implement 6 modules + cache + Serper/Gemini wiring", "Working pipeline CLI"],
            ["4. Demo & eval polish", "Streamlit UI; error analysis; romanisation OOD checks", "Screenshots + eval folder"],
            ["5. Documentation", "BCSE497J report, Review PPT, manuscript notes", "This report"],
        ],
        widths=[1.6, 2.8, 2.0],
    )

    # ------------------------------------------------------------------
    # CHAPTER 3
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 3", size=16, bold=True, center=True)
    add_heading_line(doc, "TECHNICAL SPECIFICATION", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "3.1 Requirements", size=12, bold=True)
    add_heading_line(doc, "3.1.1 Functional", size=12, bold=True, space_before=6)
    add_bullet(doc, "Ingest the 1400 CSV; map label_name to REAL/FAKE; produce train/dev/test CSVs for both splits.")
    add_bullet(doc, "Fine-tune a chosen encoder and write test metrics JSON plus prediction CSV.")
    add_bullet(doc, "Run classical TF-IDF baselines on the same splits for fair comparison.")
    add_bullet(doc, "Accept Tanglish/English claim text; return verdict, confidence, relation, explanation, citations.")
    add_bullet(doc, "Expose intermediate artefacts: normalized claim, queries, ranked evidence list.")
    add_bullet(doc, "Provide a Streamlit interface with prepared demo presets and free-text input.")

    add_heading_line(doc, "3.1.2 Non-Functional", size=12, bold=True, space_before=6)
    add_bullet(doc, "Reliability: on missing API keys, timeouts or empty evidence, fail loudly with notes — never invent sources.")
    add_bullet(doc, "Performance: Approach A inference is near-instant after training; Approach B may take tens of seconds per claim.")
    add_bullet(doc, "Cost control: disk cache for Serper/LLM prompts during repeated demos.")
    add_bullet(doc, "Security/privacy: API keys only in local .env; do not commit secrets; YouTube handles stripped in earlier LID data.")
    add_bullet(doc, "Usability: demo labels (FAKE/REAL/…) are for the presenter; the model still computes its own verdict.")
    add_bullet(doc, "Maintainability: SearchProvider interface allows swapping Serper later without rewriting the verifier.")

    add_heading_line(doc, "3.2 Feasibility Study", size=12, bold=True)
    add_heading_line(doc, "3.2.1 Technical Feasibility", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Base-sized transformers (≈110M parameters class) and 1400 training rows are feasible on "
        "CPU or Apple MPS. Hugging Face Trainer scripts already exist in the repository. Serper "
        "offers a free query tier adequate for demos; Gemini free-tier keys (including newer AQ. "
        "auth keys via google-genai) are sufficient for Project-I demonstrations. Webpage "
        "extraction is imperfect on YouTube/Facebook shells, so we added snippet fallback. Overall "
        "the stack is technically realistic for a student laptop.",
    )
    add_heading_line(doc, "3.2.2 Economic Feasibility", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "No paid cloud GPU was required for the reported seed-42 runs. Main monetary costs are "
        "optional API top-ups if free quotas are exhausted. Software dependencies are open source. "
        "Time cost is the larger investment (data cleaning, training hours, debugging extraction).",
    )
    add_heading_line(doc, "3.2.3 Social Feasibility", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "The intended use is defensive fact-checking and education. Preferring UNVERIFIED when "
        "evidence is weak reduces the risk of wrongly branding a claim as FAKE. Sources used for "
        "the 1400 set are publicly published fact-check and news pages. The demo should not be "
        "presented as an infallible authority.",
    )

    add_heading_line(doc, "3.3 System Specification", size=12, bold=True)
    add_heading_line(doc, "3.3.1 Hardware Specification", size=12, bold=True, space_before=6)
    add_caption(doc, "Table 3.1 Hardware used / recommended", above=True)
    add_table(
        doc,
        ["Item", "Specification"],
        [
            ["Development machine", "Laptop/desktop with multi-core CPU (GPU optional)"],
            ["RAM", "Recommended ≥16 GB for comfortable transformer fine-tuning"],
            ["Storage", "Several GB for model checkpoints and caches"],
            ["Network", "Required for Serper, Gemini, and page downloads in Approach B"],
        ],
        widths=[2.0, 4.3],
    )
    add_heading_line(doc, "3.3.2 Software Specification", size=12, bold=True, space_before=6)
    add_caption(doc, "Table 3.2 Software stack", above=True)
    add_table(
        doc,
        ["Layer", "Tools"],
        [
            ["Language / env", "Python 3, venv, python-dotenv"],
            ["Approach A", "PyTorch, Hugging Face transformers/datasets, scikit-learn"],
            ["Approach B", "google-genai, requests, trafilatura, BeautifulSoup, rank-bm25"],
            ["Search", "Serper.dev REST API"],
            ["UI", "Streamlit"],
            ["Data formats", "CSV, JSON/JSONL, CoNLL (LID side track)"],
        ],
        widths=[2.0, 4.3],
    )

    # ------------------------------------------------------------------
    # CHAPTER 4
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 4", size=16, bold=True, center=True)
    add_heading_line(doc, "DESIGN APPROACH AND DETAILS", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "4.1 System Architecture", size=12, bold=True)
    add_body(
        doc,
        "TruthGuard is intentionally dual, not a forced cascade. Approach A and Approach B can "
        "run independently. Approach A is appropriate when the input is a formal Tamil headline "
        "and a quick REAL/FAKE score is enough. Approach B is appropriate when the user pastes a "
        "Tanglish social claim and needs an explanation with links.",
    )
    arch = doc.add_paragraph()
    arch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = arch.add_run(
        "Approach A: Tamil headline → clean → encoder → softmax → REAL/FAKE\n\n"
        "Approach B: Tanglish claim → clean → (1) claim extract → (2) queries → (3) Serper\n"
        "→ (4) page extract → (5) rank E1..Ek → (6) Gemini verdict + cites"
    )
    _font(run, size=10, italic=True)
    add_caption(doc, "Figure 4.1 Dual-pipeline architecture of TruthGuard")
    add_body(
        doc,
        "Design principle for Approach B: information retrieval must never adjudicate truth. "
        "That keeps ranking experiments separate from verifier prompts and makes UNVERIFIED a "
        "first-class outcome when retrieval fails.",
    )

    add_heading_line(doc, "4.2 Design", size=12, bold=True)
    add_heading_line(doc, "4.2.1 Data Flow Diagram", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Context level (Approach B). External entities: End User, Serper API, Remote Websites, "
        "Gemini API. The system receives claim text and returns a structured verdict package "
        "(label, confidence, relation, explanation, evidence IDs).",
    )
    add_body(
        doc,
        "Level-1 processes. P1 Clean & extract claim; P2 Generate search queries; P3 Web search; "
        "P4 Fetch & extract page text; P5 Rank evidence; P6 Verify claim against top evidence. "
        "Data store D1: on-disk JSON cache keyed by prompt/query. Failed extraction may still "
        "pass a Serper snippet into P5/P6 so the demo does not die on JS-heavy pages.",
        first_indent=False,
    )

    add_heading_line(doc, "4.2.2 Use Case Diagram", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Primary actor: Student/reviewer using the Streamlit app. Use cases include: select a "
        "prepared demo claim; edit free-text claim; verify claim; view verdict banner; read "
        "explanation; open cited source titles; expand pipeline details (normalized claim, "
        "queries, notes). Secondary actors: Serper, Gemini, and destination websites. A separate "
        "offline use case exists for the engineer: train/evaluate Approach A models from CLI.",
    )

    add_heading_line(doc, "4.2.3 Class Diagram", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Shared dataclasses (schemas.py): ExtractedClaim, SearchResult, Evidence, RankedEvidence, "
        "Verdict, PipelineResult. LLMClient wraps Gemini generate_content with retries and JSON "
        "parsing. SearchProvider is abstract; SerperSearchProvider is the concrete class. Functions "
        "extract_from_url, rank_evidence and verify_claim operate on those dataclasses. "
        "run_pipeline orchestrates stages and records stage_reached / notes for early exits "
        "(LLM unavailable, no claim, search unavailable, no results).",
    )

    add_heading_line(doc, "4.2.4 Sequence Diagram", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Happy path: User clicks Verify → UI calls run_pipeline → LLMClient.extract claim → "
        "LLMClient.generate queries → for each query Serper.search → for each URL extract_from_url "
        "→ rank_evidence → LLMClient.verify → UI renders verdict, explanation and citations. "
        "Alternate paths: missing GEMINI_API_KEY stops at llm_unavailable; Serper failure stops at "
        "search_unavailable; zero usable evidence still calls verifier which should return "
        "UNVERIFIED with an explicit explanation.",
    )

    # ------------------------------------------------------------------
    # CHAPTER 5
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 5", size=16, bold=True, center=True)
    add_heading_line(doc, "METHODOLOGY AND TESTING", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "5.1 Module Description", size=12, bold=True)
    add_heading_line(doc, "5.1.1 Dataset construction (Approach A)", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Primary file: data/raw/tamil_news_dataset_1400_balanced.csv. Each row includes id, "
        "headline, text (article body), label/label_name, source, url and length fields. For "
        "training we mainly use the headline to reduce body-style cues. Class balance is exact: "
        "700 FAKE and 700 REAL. FAKE sources are dominated by FactCrescendoTamil, with smaller "
        "counts from BOOMTamil and NewscheckerTamil. REAL sources include Dinamalar variants, "
        "Puthiya Thalaimurai and Vikatan. All 1400 rows were retained after cleaning (UTF-8-SIG "
        "load, label mapping, clean_comment on headlines, dedupe checks).",
    )
    add_caption(doc, "Table 5.1 Source × label summary (conceptual)", above=True)
    add_table(
        doc,
        ["Source family", "Typical label", "Role in hard split"],
        [
            ["FactCrescendoTamil", "FAKE", "Train FAKE pool"],
            ["BOOM / Newschecker", "FAKE", "Test FAKE (held out)"],
            ["Dinamalar*", "REAL", "Train REAL pool"],
            ["Puthiya Thalaimurai / Vikatan", "REAL", "Test REAL (held out)"],
        ],
        widths=[2.4, 1.4, 2.5],
    )
    add_body(
        doc,
        "Hard split sizes: train 630, dev 70, test 382 (test has only 32 FAKE vs 350 REAL because "
        "held-out fact-checkers are small). Therefore we report macro-F1 and per-class behaviour, "
        "not accuracy alone. Random split (seed 42): 1120 / 140 / 140 stratified — used only as a "
        "learnability / leakage diagnostic.",
    )

    add_heading_line(doc, "5.1.2 Approach A training", size=12, bold=True, space_before=6)
    add_body(
        doc,
        "Each transformer is fine-tuned as a 2-label sequence classifier. Typical shared settings: "
        "max length 64 (headlines are short), a few epochs, batch size 16, learning rate on the "
        "order of 2e-5, AdamW, checkpoint selection by development macro-F1. Models compared: "
        "IndicBERTv2 (ai4bharat/IndicBERTv2-MLM-only), MuRIL (google/muril-base-cased), bert-base-"
        "multilingual-cased, and xlm-roberta-base. Classical baselines vectorize headlines with "
        "TF-IDF and train MultinomialNB, LogisticRegression or LinearSVC. A style-only logistic "
        "model uses a tiny handcrafted feature set (for example punctuation/question markers) to "
        "show how far shallow style cues go without semantics.",
    )

    add_heading_line(doc, "5.1.3 Approach B stages", size=12, bold=True, space_before=6)
    add_bullet(doc, "Claim extraction: Gemini detects a checkable fact and rewrites one clear English sentence.")
    add_bullet(doc, "Query generation: three diverse short queries (not the raw Tanglish string).")
    add_bullet(doc, "Web search: Serper organic results; URLs de-duplicated across queries.")
    add_bullet(doc, "Extraction: download HTML; Trafilatura then BeautifulSoup; snippet fallback on failure.")
    add_bullet(doc, "Ranking: weighted mix of TF-IDF/BM25 relevance, curated domain authority, freshness.")
    add_bullet(doc, "Verification: Gemini compares claim vs top evidence; must cite E1..Ek style IDs.")

    add_heading_line(doc, "5.2 Testing", size=12, bold=True)
    add_body(
        doc,
        "Testing mixes automated metrics and manual demo checks. For Approach A, the held-out hard "
        "test CSV is the main gate: we store Accuracy, Precision/Recall/F1 on FAKE, macro-F1, and "
        "confusion counts under eval_results_1400/. We also compared MuRIL vs mBERT predictions "
        "(they disagree on a substantial fraction of hard-test rows, so they are not identical "
        "artefacts). XLM-RoBERTa was audited because of suspiciously low Acc; it predicted FAKE for "
        "roughly two-thirds of REAL hard-test headlines. Romanisation OOD tests apply a "
        "transliteration to test text without retraining to estimate script-shift fragility.",
    )
    add_body(
        doc,
        "For Approach B, acceptance tests are scenario-based: the RBI ₹5000 preset should tend to "
        "FAKE when fact-checks are retrieved; Chandrayaan soft-landing wording should tend to REAL; "
        "over-broad holiday claims often become MISLEADING; thin WhatsApp forwards often become "
        "UNVERIFIED. We also verified that missing packages (trafilatura/bs4) previously caused "
        "total extraction failure — installing them restored page text and sensible verdicts.",
    )

    # ------------------------------------------------------------------
    # CHAPTER 6
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 6", size=16, bold=True, center=True)
    add_heading_line(doc, "PROJECT DEMONSTRATION", size=14, bold=True, center=True, space_before=4)
    add_body(
        doc,
        "Approach B is demonstrated through a Streamlit application (apps/approach_b_demo.py). "
        "From the project root, after configuring GEMINI_API_KEY and SERPER_API_KEY in .env, run:",
        first_indent=False,
    )
    code = doc.add_paragraph()
    run = code.add_run(".venv/bin/streamlit run apps/approach_b_demo.py")
    _font(run, size=10, name="Courier New")
    add_body(
        doc,
        "The sidebar/presets load Tanglish claims chosen for viva clarity. The main pane shows "
        "the claim box, Verify button, verdict banner (colour-coded), explanation paragraph, "
        "cited sources with ranking scores, and an expander for normalized claim, queries and "
        "pipeline notes. Figures 6.1–6.4 are screenshots from live runs on our machine.",
    )

    shots = [
        (IMAGES / "image.png", "Figure 6.1 FAKE verdict for the viral RBI ₹5000 note rumour (Tanglish input). "
         "Relation CONTRADICTS with citations to fact-check pages."),
        (IMAGES / "image copy.png", "Figure 6.2 REAL verdict for Chandrayaan-3 soft landing claim. "
         "Evidence supports ISRO’s announced soft landing."),
        (IMAGES / "image copy 2.png", "Figure 6.3 MISLEADING verdict for a statewide school-holiday claim. "
         "Evidence shows district-wise holidays, not necessarily all schools statewide."),
        (IMAGES / "image copy 3.png", "Figure 6.4 UNVERIFIED verdict for a free-iPhone WhatsApp rumour. "
         "Retrieved pages discuss similar scams but not a confirmed Chennai government distribution."),
    ]
    for path, cap in shots:
        if path.exists():
            add_picture(doc, path, width=5.7)
            add_caption(doc, cap)

    add_body(
        doc,
        "During demos, latency is dominated by search + page downloads + three LLM calls "
        "(claim, queries, verdict). Caching makes re-runs of the same claim much faster. If "
        "extraction fails on social platforms, snippet fallback still feeds the verifier so the "
        "UI remains usable for Review.",
    )

    # ------------------------------------------------------------------
    # CHAPTER 7
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 7", size=16, bold=True, center=True)
    add_heading_line(doc, "RESULT AND DISCUSSION", size=14, bold=True, center=True, space_before=4)

    add_heading_line(doc, "7.1 Approach A — hard split (primary)", size=12, bold=True)
    add_body(
        doc,
        "Table 7.1 summarises seed-42 results on the source-disjoint hard test. This is the table "
        "we ask evaluators to treat as the main Approach A outcome. Accuracy alone is optimistic "
        "when REAL dominates the test set; macro-F1 and qualitative error checks matter more.",
    )
    add_caption(doc, "Table 7.1 Hard-split (source-disjoint) results — seed 42", above=True)
    add_table(
        doc,
        ["Model", "Acc", "Macro-F1", "Observation"],
        [
            ["IndicBERTv2", "0.8665", "0.623", "Best transformer on this protocol"],
            ["TF-IDF + LogReg", "0.8743", "0.641", "Competitive Acc vs transformers"],
            ["TF-IDF + NB", "0.8455", "0.679", "Stronger FAKE recall in error analysis"],
            ["TF-IDF + LinearSVC", "0.8063", "0.588", "Solid classical baseline"],
            ["mBERT", "0.6126", "0.420", "Large drop vs random split"],
            ["MuRIL", "0.6099", "0.419", "Similar band to mBERT here"],
            ["XLM-RoBERTa", "0.3246", "0.266", "Collapsed; many REAL→FAKE"],
            ["Style-only LR", "0.5759", "0.402", "Style cues fail when sources change"],
        ],
        widths=[1.7, 0.9, 1.0, 2.7],
    )

    add_heading_line(doc, "7.2 Approach A — random split (diagnostic only)", size=12, bold=True)
    add_body(
        doc,
        "On the random split, IndicBERTv2, mBERT and XLM-RoBERTa reached 1.0 Acc on the 140-row "
        "test set; MuRIL reached 0.9929. These numbers show the task is learnable when the same "
        "source families appear in train and test. They do not prove real-world fake-news "
        "detection ability. We therefore refuse to lead the report with “≈100% accuracy.”",
    )
    add_caption(doc, "Table 7.2 Random-split diagnostic summary", above=True)
    add_table(
        doc,
        ["Model group", "Acc range", "How to read it"],
        [
            ["Transformers", "0.9929–1.000", "Leakage / shortcut upper bound"],
            ["TF-IDF baselines", "≈0.84–0.88", "Lower than transformers when sources overlap"],
            ["Style-only LR", "≈0.89", "Confirms shallow cues help on random split"],
        ],
        widths=[2.0, 1.4, 2.9],
    )

    add_heading_line(doc, "7.3 Error analysis and extra checks", size=12, bold=True)
    add_body(
        doc,
        "Error analysis on IndicBERTv2 hard predictions showed weaker recall on some held-out "
        "FAKE sources (for example BOOM headlines) compared with TF-IDF+NB in our sample. False "
        "negatives often lacked obvious fact-check cue words. Near-duplicate scanning found only "
        "two near-duplicate pairs under a high similarity threshold, so clone leakage is not the "
        "main issue — source alignment is. Romanisation without retraining hurt several models, "
        "reminding us that Approach A headlines are mostly native-script Tamil, not Roman Tanglish.",
    )

    add_heading_line(doc, "7.4 Approach B — qualitative results", size=12, bold=True)
    add_body(
        doc,
        "We do not invent a numeric Accuracy% for Approach B in this report. Instead, Chapter 6 "
        "screenshots document behaviour on controlled presets. In successful runs, the RBI ₹5000 "
        "rumour is labelled FAKE with fact-check citations; confirmed public science events lean "
        "REAL; partially true but overclaimed statements lean MISLEADING; underspecified forwards "
        "lean UNVERIFIED. This four-way label set is more honest for social claims than forcing "
        "every input into REAL/FAKE.",
    )

    add_heading_line(doc, "7.5 Discussion and cost notes", size=12, bold=True)
    add_body(
        doc,
        "What worked. The hard-split story is teachable and defensible in a viva: students can "
        "explain why random Acc looked “too good.” IndicBERTv2 beat other transformers under our "
        "primary protocol. Approach B’s module separation made debugging possible (for example "
        "missing extraction libraries looked like “no evidence” until dependencies were installed).",
    )
    add_body(
        doc,
        "What was harder than expected. XLM-RoBERTa’s hard-split failure shows that a famous "
        "multilingual model is not automatically best for this small Tamil set. Hard-test FAKE "
        "count is small, so FAKE F1 stays noisy. Approach B depends on third-party APIs and site "
        "HTML quality; live demos need keys and network. Ranking weights (relevance/authority/"
        "freshness) are reasoned defaults, not exhaustively tuned.",
    )
    add_body(
        doc,
        "Cost analysis (as applicable). Approach A: one-time training cost (CPU hours) then nearly "
        "free inference. Approach B: per-claim cost ≈ several LLM calls + search queries + HTTP "
        "fetches; mitigated by caching during rehearsal. No separate cloud bill was required for "
        "the numbers reported here beyond optional free-tier APIs.",
    )

    # ------------------------------------------------------------------
    # CHAPTER 8
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 8", size=16, bold=True, center=True)
    add_heading_line(doc, "CONCLUSION", size=14, bold=True, center=True, space_before=4)
    add_body(
        doc,
        "TruthGuard meets the Project-I goal of a concrete, explainable system for Tamil/Tanglish "
        "misinformation work. Approach A contributes a self-collected 1400-headline corpus, multiple "
        "model comparisons, and an evaluation protocol that refuses to hide behind random-split "
        "ceilings. Approach B contributes a working evidence pipeline and Streamlit demonstration "
        "with citeable verdicts. The dual design clarifies a practical trade-off: classifiers are "
        "cheap and fast on formal headlines; retrieval-augmented checking is slower but better "
        "matched to fresh Tanglish rumours.",
    )
    add_body(
        doc,
        "Future work. (1) Finish and report multi-seed mean±std if additional seeds complete. "
        "(2) Label a small Tanglish claim set so Approach B can quote Acc/F1 honestly. "
        "(3) Increase FAKE source diversity for a healthier hard-test balance. "
        "(4) Add romanisation / code-mix augmentation for Approach A. "
        "(5) Optional batch runner for Approach B with rate limits. "
        "(6) Draw formal UML diagrams in a diagramming tool for the final hardbound copy if required.",
    )

    add_heading_line(doc, "Individual Contributions", size=12, bold=True)
    add_body(
        doc,
        "All members reviewed results and the demo together. The ownership split below reflects "
        "primary responsibility areas for Project-I / Review documentation.",
        first_indent=False,
    )
    add_caption(doc, "Table 8.1 Team contribution split", above=True)
    add_table(
        doc,
        ["Name", "Reg. No.", "Main contributions"],
        [
            [
                "Pranav Kalra",
                "23BCE0504",
                "TF-IDF and style baselines; error analysis; romanisation OOD checks; eval packaging; PPT support",
            ],
            [
                "Samya Mehta",
                "23BDS0204",
                "Approach B modules and caching; Streamlit demo; Gemini client updates; report drafting",
            ],
            [
                "Palak Goyal",
                "23BDS0209",
                "1400 dataset curation/cleaning; Approach A splits; transformer training runs; metrics tables",
            ],
        ],
        widths=[1.4, 1.2, 3.8],
    )

    # ------------------------------------------------------------------
    # CHAPTER 9
    # ------------------------------------------------------------------
    add_heading_line(doc, "CHAPTER 9", size=16, bold=True, center=True)
    add_heading_line(doc, "REFERENCES", size=14, bold=True, center=True, space_before=4)
    refs = [
        "R. K. Kaliyar, A. Goswami, and P. Narang, “FakeBERT: Fake news detection in social media with a BERT-based deep learning approach,” Multimedia Tools and Applications, 2021.",
        "S. Khanuja et al., “MuRIL: Multilingual Representations for Indian Languages,” arXiv:2103.10730, 2021.",
        "D. Kakwani et al., “IndicNLPSuite: Monolingual Corpora, Evaluation Benchmarks and Pre-trained Multilingual Language Models for Indian Languages,” Findings of EMNLP, 2020.",
        "S. Doddapaneni et al., “Towards Leaving No Indic Language Behind: Building Monolingual Corpora, Benchmark and Models for Indic Languages,” ACL, 2023 (IndicBERTv2).",
        "J. Thorne, A. Vlachos, C. Christodoulopoulos, and A. Mittal, “FEVER: a Large-scale Dataset for Fact Extraction and VERification,” NAACL, 2018.",
        "P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” NeurIPS, 2020.",
        "B. R. Chakravarthi et al., DravidianLangTech shared tasks on code-mixed Dravidian language processing.",
        "Google AI, Gemini API documentation. [Online]. Available: https://ai.google.dev/",
        "Serper.dev, Google Search API documentation. [Online]. Available: https://serper.dev/",
        "Project evaluation notes, eval_results_1400/RESULTS.md (internal technical summary).",
    ]
    for i, ref in enumerate(refs, 1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.3)
        para.paragraph_format.first_line_indent = Inches(-0.3)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(f"[{i}] {ref}")
        _font(run, size=11)

    # APPENDIX
    add_heading_line(doc, "APPENDIX A – SAMPLE CODE", size=14, bold=True, center=True)
    add_body(
        doc,
        "This appendix lists only entry points. Full source lives in the project repository under "
        "src/fake_news_classifier/, src/verification/ and apps/.",
        first_indent=False,
    )
    add_body(doc, "A.1 Approach B demo", first_indent=False)
    code = doc.add_paragraph()
    run = code.add_run(
        "# from project root, with .env configured\n"
        ".venv/bin/streamlit run apps/approach_b_demo.py\n\n"
        "# CLI single-claim run\n"
        ".venv/bin/python src/verification/run_demo.py --text "
        "\"RBI 5000 rupees note release pannitanga nu news la varudhu\""
    )
    _font(run, size=9, name="Courier New")
    add_body(doc, "A.2 Approach A training (example)", first_indent=False)
    code2 = doc.add_paragraph()
    run = code2.add_run(
        ".venv/bin/python src/fake_news_classifier/preprocess_1400.py --split hard\n"
        ".venv/bin/python src/fake_news_classifier/train_classifier.py --model indicbert-v2\n"
        "# metrics written under models_fake_news_1400_hard/ or eval_results_1400/"
    )
    _font(run, size=9, name="Courier New")
    add_body(
        doc,
        "A.3 Orchestrator. The function run_pipeline(raw_text) in src/verification/pipeline.py "
        "wires claim extraction → queries → search → extraction → ranking → verification and "
        "returns a PipelineResult with every intermediate field for debugging and viva explanation.",
        first_indent=False,
    )
